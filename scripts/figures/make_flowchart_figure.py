"""Fig. 3 of the paper: the three-stage flowchart, corrected on its original design.

Starts from results/training_pipeline_flowchart_editable.pptx (the W1/W2/W3 flowchart) and
changes only what was wrong, keeping its shapes, colours and layout:

  * COCO weights (yolov5s.pt), not ImageNet.
  * The paper's dataset names; parentheses give nc where it is not 11 (data.yaml).
  * Stage II pools datasets and trains from COCO. Its output used later is the pooled
    datasets, so its cylinder is the pooled data (D2), and there is no arrow from W1 into it.
  * Stage III starts from W1, the IR-Pot-Background Stage I checkpoint, so a connector runs
    from W1 to Stage III.

Writes results/fig3_training_flowchart.pptx, renders it with LibreOffice and saves
results/ieee/IEEE_LaTeX/figs/fig3_three_stage_training.png (300 dpi, cropped).

Usage (from the project root, with python-pptx installed, e.g. venv/bin/python):
  python scripts/figures/make_flowchart_figure.py
"""

import copy
import subprocess
import tempfile
from pathlib import Path

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "results" / "training_pipeline_flowchart_editable.pptx"
OUT_PPTX = ROOT / "results" / "fig3_training_flowchart.pptx"
OUT_PNG = ROOT / "results" / "ieee" / "IEEE_LaTeX" / "figs" / "fig3_three_stage_training.png"
BLUE = RGBColor(0x2E, 0x75, 0xB6)
GREY = RGBColor(0x59, 0x59, 0x59)
DATA_PT = 12   # dataset lines; the original's 18 pt short names do not fit the paper's names

STAGES = {
    7: ("Stage I", "Multi-Spectral Base Training", [
        "UV-Soil-Only-Baseline · UV-Soil-Only-Target",
        "IR-Soil-Only (7c) · UV-Strawberry-Pot (9c)",
        "IR-Pot-Background (5c) · UV-Pot-Background (5c)",
        "Mixed Color and Background (5c)"]),
    8: ("Stage II", "Pooled Training (COCO init)", [
        "UV-Combined: UV-Soil-Only-Baseline",
        "+ UV-Strawberry-Pot",
        "IR-Combined (5c): IR-Pot-Background",
        "+ UV-Pot-Background"]),
    9: ("Stage III", "Domain-Specific Fine-Tuning", [
        "UV-Combined · IR-Combined (5c)",
        "IR-Soil-Only (7c) · UV-Soil-Only-Target",
        "Mixed Color and Background (5c)"]),
}
TEXT = {
    2: ["COCO Pre-trained Weights (yolov5s.pt)"],
    16: ["Best weights", "IR-Pot-Background", "best.pt"],
    17: ["Pooled", "datasets"],
    6: ["Final output - detected moisture levels"],
}


def set_runs(par, text):
    """Keep the paragraph's first run (its formatting) and give it all the text."""
    runs = par.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_text(shape, lines):
    pars = shape.text_frame.paragraphs
    for _ in range(len(lines) - len(pars)):      # more lines than paragraphs: copy the last
        pars[-1]._p.addnext(copy.deepcopy(pars[-1]._p))
    pars = shape.text_frame.paragraphs
    for par, line in zip(pars, lines):
        set_runs(par, line)
    for par in pars[len(lines):]:
        par._p.getparent().remove(par._p)


def set_stage(shape, title, subtitle, data):
    pars = shape.text_frame.paragraphs
    set_runs(pars[0], title)
    set_runs(pars[1], subtitle)
    template = pars[3]._p            # first dataset line: grey, not bold
    for par in pars[3:]:
        par._p.getparent().remove(par._p)
    body = shape.text_frame._txBody
    for line in data:
        p = copy.deepcopy(template)
        body.append(p)
    for par, line in zip(shape.text_frame.paragraphs[3:], data):
        set_runs(par, line)
        for r in par.runs:
            r.font.size, r.font.bold = Pt(DATA_PT), False


def connector(slide, x1, y1, x2, y2, head=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb, c.line.width = BLUE, Pt(6)
    if head:
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                              {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return c


def main():
    prs = Presentation(SRC)
    slide = prs.slides[0]
    shapes = {s.shape_id: s for s in slide.shapes}

    for sid, lines in TEXT.items():
        set_text(shapes[sid], lines)
    shapes[16].left, shapes[16].width = Inches(10.68), Inches(2.6)
    for par in shapes[16].text_frame.paragraphs:
        for r in par.runs:
            r.font.size = Pt(18)
    shapes[6].width = Inches(6.0)                 # the original wraps this line in two
    for sid, (title, subtitle, data) in STAGES.items():
        set_stage(shapes[sid], title, subtitle, data)
    shapes[7].left, shapes[7].width = Inches(4.33), Inches(4.0)   # widest dataset lines
    shapes[7].height = Inches(2.2)
    shapes[8].height = Inches(2.1)
    shapes[2].left, shapes[2].width = Inches(1.83), Inches(9.0)    # title on one line
    shapes[20].width = Inches(0.62)               # images -> Stage I arrow, shortened to fit
    shapes[21].left, shapes[21].width = Inches(8.4), Inches(0.9)  # Stage I -> W1
    set_text(shapes[12], ["D2"])                  # Stage II's output used downstream is data

    # W1 does not feed Stage II (it starts from COCO): drop that arrow and route W1 to Stage III.
    a = shapes[22]
    a._element.getparent().remove(a._element)
    x_w1 = 9.955                                  # W1 centre
    x_side, y_low = 12.45, 7.16                   # right of Stage II, level with Stage III
    st3 = shapes[9]
    x_end = (st3.left + st3.width) / 914400 + 0.05
    connector(slide, x_w1, 2.75, x_w1, 3.30)
    connector(slide, x_w1 - 0.04, 3.30, x_side + 0.04, 3.30)
    connector(slide, x_side, 3.26, x_side, y_low + 0.04)
    connector(slide, x_side + 0.04, y_low, x_end, y_low, head=True)
    tb = slide.shapes.add_textbox(Inches(8.0), Inches(7.25), Inches(4.3), Inches(0.4))
    tb.text_frame.text = "Stage III starts from W1"
    run = tb.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.name, run.font.italic = Pt(18), "Arial", True
    run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)

    prs.save(OUT_PPTX)

    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(OUT_PPTX)],
                       check=True, capture_output=True)
        pdf = Path(tmp) / (OUT_PPTX.stem + ".pdf")
        subprocess.run(["pdftoppm", "-r", "300", "-png", "-singlefile", str(pdf), str(Path(tmp) / "f")],
                       check=True)
        im = Image.open(Path(tmp) / "f.png").convert("RGB")
    bbox = ImageChops.difference(im, Image.new("RGB", im.size, "white")).getbbox()
    pad = 30
    im.crop((max(0, bbox[0] - pad), max(0, bbox[1] - pad),
             min(im.width, bbox[2] + pad), min(im.height, bbox[3] + pad))).save(OUT_PNG)
    print("saved", OUT_PPTX, OUT_PNG)


if __name__ == "__main__":
    main()
